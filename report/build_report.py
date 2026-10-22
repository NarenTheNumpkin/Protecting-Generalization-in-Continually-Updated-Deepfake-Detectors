"""Build the project report (Word) from the VIT template and the content files.

Usage: python build_report.py <template.docx> <out.docx> [pages.json]
pages.json (optional) holds page numbers measured on a previous render, used for the cached
contents, list of figures and list of tables (Word recomputes them when the fields update)."""

import copy
import json
import os
import re
import sys
import zipfile

import docx
import pypandoc
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from refs import REFS, apa_entry, cite, citet, sort_key  # noqa: E402

TEMPLATE, OUT = sys.argv[1], sys.argv[2]
PAGES = json.load(open(sys.argv[3])) if len(sys.argv) > 3 and os.path.exists(sys.argv[3]) else {}

TITLE = "Protecting Generalization in Continually Updated Deepfake Detectors"
TEXT_W = 5.77  # inches: A4 width 8.27 minus 1.5 left and 1.0 right margins
MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
CHAPTER_FILES = ["ch1.txt", "ch2.txt", "ch3.txt", "ch4.txt", "ch5.txt", "ch6.txt"]

# ======================================================================================= parsing
def parse_file(path):
    """Return a list of blocks: (kind, data)."""
    lines = open(path).read().split("\n")
    blocks, i, para = [], 0, []

    def flush():
        if para:
            blocks.append(("p", " ".join(x.strip() for x in para)))
            para.clear()

    def lab(s):
        m = re.search(r"\{((?:ch|sec|app):[A-Za-z0-9_-]+)\}\s*$", s)
        return (s[: m.start()].strip(), m.group(1)) if m else (s.strip(), None)

    while i < len(lines):
        ln = lines[i]
        if not ln.strip():
            flush(); i += 1; continue
        if ln.startswith("@chapter ") or ln.startswith("@appendix ") or ln.startswith("@section ") or ln.startswith("@subsection "):
            flush()
            kind, rest = ln.split(" ", 1)
            title, label = lab(rest)
            blocks.append((kind[1:], {"title": title, "label": label}))
        elif ln.startswith("@fig "):
            flush()
            path, width, label, caption = [x.strip() for x in ln[5:].split("|", 3)]
            blocks.append(("fig", {"path": path, "width": float(width), "label": label, "caption": caption}))
        elif ln.startswith("@eq "):
            flush()
            label, tex = [x.strip() for x in ln[4:].split("|", 1)]
            blocks.append(("eq", {"label": label, "tex": tex}))
        elif ln.startswith("@note "):
            flush()
            blocks.append(("note", ln[6:].strip()))
        elif ln.startswith("@table "):
            flush()
            parts = [x.strip() for x in ln[7:].split("|")]
            label, caption = parts[0], parts[1]
            opts = dict(p.split("=", 1) for p in parts[2:])
            rows = []
            i += 1
            while lines[i].strip() != "@endtable":
                rows.append([c.strip() for c in lines[i].split(" | ")])
                i += 1
            blocks.append(("table", {"label": label, "caption": caption, "rows": rows,
                                     "widths": [float(w) for w in opts["widths"].split(",")], "align": opts.get("align", ""),
                                     "size": float(opts.get("size", 10)), "intervals": opts.get("intervals") == "1"}))
        elif ln.startswith("@alg "):
            flush()
            label, caption = [x.strip() for x in ln[5:].split("|", 1)]
            req, steps = None, []
            i += 1
            while lines[i].strip() != "@endalg":
                l2 = lines[i]
                if l2.startswith("@require "):
                    req = l2[9:].strip()
                elif l2.strip():
                    lvl, txt = l2.split("|", 1)
                    steps.append((int(lvl), txt.strip()))
                i += 1
            blocks.append(("alg", {"label": label, "caption": caption, "require": req, "steps": steps}))
        elif ln.startswith("@code"):
            flush()
            code = []
            i += 1
            while lines[i].strip() != "@endcode":
                code.append(lines[i])
                i += 1
            blocks.append(("code", "\n".join(code)))
        elif ln.startswith("@"):
            raise ValueError("unknown directive: " + ln)
        else:
            para.append(ln)
        i += 1
    flush()
    return blocks


CONTENT = os.path.join(HERE, "content")
chapters = [parse_file(os.path.join(CONTENT, f)) for f in CHAPTER_FILES]
appendix = parse_file(os.path.join(CONTENT, "appendix.txt"))

# ======================================================================================= numbering
LABELS = {}  # label -> number string
heads = []   # (level, number, title, label) for TOC
figs, tabs = [], []  # (number, caption)


def number_blocks(blocks, prefix, chap_no, is_appendix=False):
    sec = sub = fig = tab = eq = alg = 0
    for kind, d in blocks:
        if kind in ("chapter", "appendix"):
            if d["label"]:
                LABELS[d["label"]] = str(chap_no)
        elif kind == "section":
            sec += 1; sub = 0
            d["num"] = f"{chap_no}.{sec}"
            if d["label"]:
                LABELS[d["label"]] = d["num"]
        elif kind == "subsection":
            sub += 1
            d["num"] = f"{chap_no}.{sec}.{sub}"
            if d["label"]:
                LABELS[d["label"]] = d["num"]
        elif kind == "fig":
            fig += 1
            d["num"] = f"{prefix}{chap_no}.{fig}"
            LABELS[d["label"]] = d["num"]
        elif kind == "table":
            tab += 1
            d["num"] = f"{prefix}{chap_no}.{tab}"
            LABELS[d["label"]] = d["num"]
        elif kind == "eq":
            eq += 1
            d["num"] = f"{chap_no}.{eq}"
            LABELS[d["label"]] = d["num"]
        elif kind == "alg":
            alg += 1
            d["num"] = f"{chap_no}.{alg}"
            LABELS[d["label"]] = d["num"]


for n, blocks in enumerate(chapters, 1):
    number_blocks(blocks, "", n)
app_no = 0
for kind, d in appendix:
    if kind == "appendix":
        app_no += 1
        d["num"] = str(app_no)
        if d["label"]:
            LABELS[d["label"]] = str(app_no)
# appendix tables are numbered A<k>.<i>
cur, tcount = 0, 0
for kind, d in appendix:
    if kind == "appendix":
        cur += 1; tcount = 0
    elif kind == "table":
        tcount += 1
        d["num"] = f"A{cur}.{tcount}"
        LABELS[d["label"]] = d["num"]
    elif kind in ("fig", "eq", "alg", "section", "subsection"):
        raise ValueError("appendix supports tables, notes, paragraphs and code only")

USED_CITES = set()


def resolve_refs(text):
    def rep(m):
        lab = m.group(1)
        if lab not in LABELS:
            raise KeyError("undefined label " + lab)
        n = LABELS[lab]
        if lab.startswith("fig:"):
            return "Figure " + n
        if lab.startswith("eq:"):
            return f"({n})"
        return n
    text = re.sub(r"\{((?:fig|tab|eq|sec|alg|ch|app):[A-Za-z0-9_-]+)\}", rep, text)

    def c1(m):
        keys = [k.strip() for k in m.group(1).split(",")]
        for k in keys:
            if k not in REFS:
                raise KeyError("missing reference " + k)
        USED_CITES.update(keys)
        return cite(keys)

    def c2(m):
        k = m.group(1).strip()
        if k not in REFS:
            raise KeyError("missing reference " + k)
        USED_CITES.add(k)
        return citet(k)
    text = re.sub(r"\\citet\{([^}]+)\}", c2, text)
    text = re.sub(r"\\cite\{([^}]+)\}", c1, text)
    return text


# ======================================================================================= math (LaTeX -> OMML via pandoc)
MATH = {}  # tex -> oMath element (inline form)


def collect_math(s):
    for m in re.finditer(r"(?<!\\)\$(.+?)(?<!\\)\$", s):
        MATH.setdefault(m.group(1), None)


def walk_text(blocks):
    for kind, d in blocks:
        if kind in ("p", "note"):
            collect_math(d)
        elif kind == "table":
            collect_math(d["caption"])
            for r in d["rows"]:
                for c in r:
                    collect_math(c)
        elif kind == "fig":
            collect_math(d["caption"])
        elif kind == "eq":
            MATH.setdefault("DISPLAY::" + d["tex"], None)
        elif kind == "alg":
            collect_math(d["caption"]); collect_math(d["require"] or "")
            for _, t in d["steps"]:
                collect_math(t)


for b in chapters:
    walk_text(b)
walk_text(appendix)
SYMBOLS = [
    (r"\theta", "trainable parameters of the detector (layer-normalization affine parameters and head)"),
    (r"\bar\theta,\ \theta_{t-1}", "parameters of the frozen copy of the model before an update"),
    (r"\mathcal{T}_t", "task $t$ of the stream; $\\mathcal{T}_1$ is the base task"),
    (r"K", "number of tasks in the stream, including the base task"),
    (r"\mathcal{D}_t", "few-shot training sample of task $t$"),
    (r"A_{k,j}", "video AUROC on the test split of task $j$ after learning task $k$"),
    (r"z,\ \bar z", "unit-length feature of a face under the current and the frozen model, $z\\in\\mathbb{S}^{1023}$"),
    (r"\ell,\ \bar\ell", "logits of the current and the frozen model"),
    (r"\mathcal{B},\ \mathcal{R}", "samples of a minibatch and its genuine samples"),
    (r"\lambda_g,\ \lambda_l", "weights of the alignment and logit-distillation terms"),
    (r"T", "temperature of the logit distillation"),
    (r"\sigma", "softmax function"),
    (r"U", "unseen-panel score (Equation {eq:U})"),
    (r"U_{\mathrm{deg}}", "mean unseen-panel score over the 13 image degradations"),
    (r"s_\theta(x)", "decision score of input $x$, the difference of the two logits"),
    (r"\delta(x)", "angular drift of the feature of $x$ caused by an update"),
    (r"\epsilon_G,\ \epsilon_Q", "bounds on the score shift of genuine faces and of unseen manipulated faces"),
]
for s, desc in SYMBOLS:
    MATH.setdefault(s, None)
    collect_math(desc)


def fix_abs(tex):
    """Turn |x| (and \\big|x\\big|) into \\left|x\\right| so that Word and LibreOffice draw real delimiters."""
    tex = tex.replace(r"\big|", "|")
    out, opened, i = [], False, 0
    while i < len(tex):
        if tex[i] == "\\" and i + 1 < len(tex):
            out.append(tex[i:i + 2]); i += 2; continue
        if tex[i] == "|":
            out.append(r"\right|" if opened else r"\left|"); opened = not opened
        else:
            out.append(tex[i])
        i += 1
    return "".join(out) if not opened else tex


def convert_math():
    keys = list(MATH.keys())
    md = []
    for i, k in enumerate(keys):
        if k.startswith("DISPLAY::"):
            md.append(f"MK{i}Q\n\n$${fix_abs(k[9:])}$$\n")
        else:
            md.append(f"MK{i}Q ${fix_abs(k)}$\n")
    out = os.path.join(HERE, "_math.docx")
    pypandoc.convert_text("\n".join(md), "docx", format="markdown", outputfile=out)
    xml = zipfile.ZipFile(out).read("word/document.xml")
    root = etree.fromstring(xml)
    body = root.find(qn("w:body"))
    cur = None
    for p in body.iter(qn("w:p")):
        txt = "".join(t.text or "" for t in p.iter(qn("w:t")))
        m = re.match(r"MK(\d+)Q", txt)
        if m:
            cur = int(m.group(1))
        om = p.find(".//{%s}oMath" % MATH_NS)
        if om is not None and cur is not None and MATH[keys[cur]] is None:
            MATH[keys[cur]] = om
    for k in keys:
        if MATH[k] is not None:
            for sc in MATH[k].iter("{%s}sepChr" % MATH_NS):
                if sc.get("{%s}val" % MATH_NS) == "":
                    sc.getparent().remove(sc)
    missing = [k for k in keys if MATH[k] is None]
    if missing:
        raise RuntimeError("math conversion failed for: " + repr(missing[:5]))


convert_math()


def math_element(tex, size_pt):
    om = copy.deepcopy(MATH[tex])
    for r in om.iter("{%s}r" % MATH_NS):
        rpr = r.find(qn("w:rPr"))
        if rpr is None:
            rpr = etree.SubElement(r, qn("w:rPr"))
            # w:rPr must follow m:rPr inside m:r
            mrpr = r.find("{%s}rPr" % MATH_NS)
            r.remove(rpr)
            r.insert(1 if mrpr is not None else 0, rpr)
        f = etree.SubElement(rpr, qn("w:rFonts"))
        for a in ("ascii", "hAnsi"):
            f.set(qn("w:" + a), "Cambria Math")
        sz = etree.SubElement(rpr, qn("w:sz")); sz.set(qn("w:val"), str(int(size_pt * 2)))
    return om


# ======================================================================================= document helpers
doc = docx.Document(TEMPLATE)
body = doc.element.body


def set_cell_borders(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders"); tcPr.append(b)
    for edge, (val, sz) in kw.items():
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), val); e.set(qn("w:sz"), str(sz)); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000")
        b.append(e)


def no_table_borders(tbl):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge); e.set(qn("w:val"), "nil"); b.append(e)
    tblPr.append(b)


def set_widths(tbl, widths):
    tbl.autofit = False
    tblPr = tbl._tbl.tblPr
    if tblPr.find(qn("w:tblLayout")) is None:
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    tw = tblPr.find(qn("w:tblW"))
    if tw is None:
        tw = OxmlElement("w:tblW"); tblPr.append(tw)
    tw.set(qn("w:w"), str(int(sum(widths) * 1440))); tw.set(qn("w:type"), "dxa")
    grid = tbl._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w * 1440)))
    for row in tbl.rows:
        for c, w in zip(row.cells, widths):
            c.width = Inches(w)


def cell_margins(tbl, lr=60, tb=20):
    tblPr = tbl._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for side, v in (("top", tb), ("left", lr), ("bottom", tb), ("right", lr)):
        e = OxmlElement("w:" + side); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    tblPr.append(m)


def para_fmt(p, align=None, first=None, left=None, hanging=None, before=None, after=None, line=None, keep_next=False,
             keep_lines=False, page_break_before=False):
    pf = p.paragraph_format
    if align is not None:
        p.alignment = align
    if first is not None:
        pf.first_line_indent = Inches(first)
    if left is not None:
        pf.left_indent = Inches(left)
    if hanging is not None:
        pf.first_line_indent = Inches(-hanging)
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = line
    pf.keep_with_next = keep_next
    if keep_lines:
        pf.keep_together = True
    if page_break_before:
        pf.page_break_before = True
    pf.widow_control = True


def add_run(p, text, size=12, bold=False, italic=False, font=None, highlight=False, color=None, caps=False):
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = font or "Times New Roman"
    rpr = r._r.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("ascii", "hAnsi", "cs", "eastAsia"):
        rf.set(qn("w:" + a), font or "Times New Roman")
    if highlight:
        h = OxmlElement("w:highlight"); h.set(qn("w:val"), "yellow"); rpr.append(h)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    if caps:
        r.font.all_caps = True
    return r


TOKEN = re.compile(r"(?<!\\)\$(.+?)(?<!\\)\$|\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`|\[\[(.+?)\]\]")


def render_inline(p, text, size=12, bold=False, italic=False):
    """Inline markup: $math$, **bold**, *italic*, `code`, [[placeholder]]; refs and cites resolved first."""
    parts = re.split(r"((?<!\\)\$.+?(?<!\\)\$)", text)
    text = "".join(x if x.startswith("$") else resolve_refs(x) for x in parts)
    pos = 0
    for m in TOKEN.finditer(text):
        if m.start() > pos:
            add_run(p, text[pos: m.start()], size, bold, italic)
        if m.group(1) is not None:
            p._p.append(math_element(m.group(1), size))
        elif m.group(2) is not None:
            render_inline(p, m.group(2), size, True, italic)
        elif m.group(3) is not None:
            render_inline(p, m.group(3), size, bold, True)
        elif m.group(4) is not None:
            add_run(p, m.group(4), size - 1, bold, italic, font="Courier New")
        elif m.group(5) is not None:
            add_run(p, "[" + m.group(5) + "]", size, bold, italic, highlight=True)
        pos = m.end()
    if pos < len(text):
        add_run(p, text[pos:], size, bold, italic)


def body_par(text, first=0.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12):
    p = doc.add_paragraph(style="Report Body")
    render_inline(p, text, size)
    return p


# ======================================================================================= styles
styles_el = doc.styles.element


def make_style(sid, name, based="Normal", size=12, bold=False, italic=False, caps=False, align=None, before=0, after=0,
               line=None, line_rule="auto", first=None, keep_next=False, outline=None, page_break=False, tabs=None,
               left=None, hanging=None):
    old = styles_el.find(qn("w:style") + f"[@{qn('w:styleId')}='{sid}']")
    if old is not None:
        styles_el.remove(old)
    s = OxmlElement("w:style"); s.set(qn("w:type"), "paragraph"); s.set(qn("w:styleId"), sid)
    if sid == "Normal":
        s.set(qn("w:default"), "1")
    n = OxmlElement("w:name"); n.set(qn("w:val"), name); s.append(n)
    if based and sid != "Normal":
        b = OxmlElement("w:basedOn"); b.set(qn("w:val"), based); s.append(b)
    q = OxmlElement("w:qFormat"); s.append(q)
    ppr = OxmlElement("w:pPr")
    if keep_next:
        ppr.append(OxmlElement("w:keepNext"))
    if page_break:
        ppr.append(OxmlElement("w:pageBreakBefore"))
    wc = OxmlElement("w:widowControl"); ppr.append(wc)
    if tabs:
        ts = OxmlElement("w:tabs")
        for kind, pos, leader in tabs:
            t = OxmlElement("w:tab"); t.set(qn("w:val"), kind); t.set(qn("w:pos"), str(pos))
            if leader:
                t.set(qn("w:leader"), leader)
            ts.append(t)
        ppr.append(ts)
    sp = OxmlElement("w:spacing"); sp.set(qn("w:before"), str(int(before * 20))); sp.set(qn("w:after"), str(int(after * 20)))
    if line:
        sp.set(qn("w:line"), str(int(line * 240))); sp.set(qn("w:lineRule"), line_rule)
    ppr.append(sp)
    if first is not None or left is not None or hanging is not None:
        ind = OxmlElement("w:ind")
        if left is not None:
            ind.set(qn("w:left"), str(int(left * 1440)))
        if first is not None:
            ind.set(qn("w:firstLine"), str(int(first * 1440)))
        if hanging is not None:
            ind.set(qn("w:hanging"), str(int(hanging * 1440)))
        ppr.append(ind)
    if align:
        jc = OxmlElement("w:jc"); jc.set(qn("w:val"), align); ppr.append(jc)
    if outline is not None:
        ol = OxmlElement("w:outlineLvl"); ol.set(qn("w:val"), str(outline)); ppr.append(ol)
    s.append(ppr)
    rpr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    for a in ("ascii", "hAnsi", "cs", "eastAsia"):
        rf.set(qn("w:" + a), "Times New Roman")
    rpr.append(rf)
    if bold:
        rpr.append(OxmlElement("w:b"))
    if italic:
        rpr.append(OxmlElement("w:i"))
    if caps:
        rpr.append(OxmlElement("w:caps"))
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size * 2))); rpr.append(sz)
    szc = OxmlElement("w:szCs"); szc.set(qn("w:val"), str(int(size * 2))); rpr.append(szc)
    s.append(rpr)
    styles_el.append(s)


# convert the cover-page heading paragraphs to explicit formatting before the heading styles are redefined
for p in doc.paragraphs:
    if p.style.name in ("Heading 1", "Heading 2", "Heading 3"):
        sty = p.style
        size = sty.font.size.pt if sty.font.size else 12
        for r in p.runs:
            if r.font.size is None:
                r.font.size = Pt(size)
            r.font.bold = True
        if p.alignment is None:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.style = doc.styles["Normal"]

RIGHT = int(TEXT_W * 1440)
make_style("Heading1", "heading 1", size=16, bold=True, align="center", before=0, after=18, line=1.15, keep_next=True,
           outline=0, page_break=True)
make_style("Heading2", "heading 2", size=14, caps=True, before=12, after=6, line=1.15, keep_next=True, outline=1)
make_style("Heading3", "heading 3", size=12, caps=True, before=10, after=4, line=1.15, keep_next=True, outline=2)
make_style("ReportBody", "Report Body", size=12, align="both", after=6, line=1.5, first=0.5)
make_style("ReportPlain", "Report Plain", size=12, align="both", after=6, line=1.5)
make_style("FigCaption", "Fig Caption", size=11, align="center", before=4, after=12, line=1.0)
make_style("TabCaption", "Tab Caption", size=11, align="center", before=10, after=4, line=1.0, keep_next=True)
make_style("TableNote", "Table Note", size=10, align="both", before=3, after=10, line=1.0)
make_style("FigurePara", "Figure Para", size=12, align="center", before=8, after=2, line=1.0, keep_next=True)
make_style("TableText", "Table Text", size=10, line=1.0)
make_style("TOC1", "toc 1", size=12, bold=True, before=8, after=2, line=1.15, tabs=[("right", RIGHT, "dot")])
make_style("TOC2", "toc 2", size=12, before=0, after=2, line=1.15, left=0.3, tabs=[("right", RIGHT, "dot")])
make_style("TOC3", "toc 3", size=12, before=0, after=2, line=1.15, left=0.6, tabs=[("right", RIGHT, "dot")])
make_style("TOC9", "toc 9", size=12, before=0, after=4, line=1.15, left=1.0, hanging=1.0, tabs=[("right", RIGHT, "dot")])
make_style("RefEntry", "Reference Entry", size=12, align="left", after=6, line=1.5, left=0.5, hanging=0.5)
make_style("CodeText", "Code Text", size=9, line=1.0)

# ======================================================================================= template front pages
P = lambda: list(body.iterchildren())  # noqa: E731


def ptext(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def has_sect(el):
    ppr = el.find(qn("w:pPr"))
    return ppr is not None and ppr.find(qn("w:sectPr")) is not None


# 1. remove the instruction pages (everything before "A project report on")
children = P()
start = next(i for i, el in enumerate(children) if el.tag == qn("w:p") and ptext(el).strip() == "A project report on")
for el in children[:start]:
    body.remove(el)
# 2. remove everything from CONTENTS to the final sectPr
children = P()
ci = next(i for i, el in enumerate(children) if el.tag == qn("w:p") and ptext(el).strip() == "CONTENTS")
for el in children[ci:]:
    if el.tag != qn("w:sectPr"):
        body.remove(el)


def find_par(substr, nth=0):
    hits = [p for p in doc.paragraphs if substr in p.text]
    return hits[nth] if len(hits) > nth else None


def set_par_text(p, segments, size=None, bold=None, italic=None, align=None):
    """Replace the runs of p by segments [(text, highlight)] keeping the first run's formatting."""
    runs = p.runs
    proto = copy.deepcopy(runs[0]._r) if runs else None
    for r in runs:
        r._r.getparent().remove(r._r)
    for text, hl in segments:
        if proto is not None:
            r = copy.deepcopy(proto)
            for x in list(r):
                if x.tag != qn("w:rPr"):
                    r.remove(x)
            rpr = r.find(qn("w:rPr"))
            if rpr is not None:
                for h in rpr.findall(qn("w:highlight")):
                    rpr.remove(h)
                for c in rpr.findall(qn("w:color")):
                    rpr.remove(c)
            for k, piece in enumerate(text.split("\t")):
                if k:
                    r.append(OxmlElement("w:tab"))
                if piece:
                    t = OxmlElement("w:t"); t.text = piece; t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                    r.append(t)
            p._p.append(r)
            from docx.text.run import Run
            run = Run(r, p)
        else:
            run = p.add_run(text)
        if size:
            run.font.size = Pt(size)
        if bold is not None:
            run.font.bold = bold
        if italic is not None:
            run.font.italic = italic
        run.font.name = "Times New Roman"
        if hl:
            rpr = run._r.get_or_add_rPr()
            h = OxmlElement("w:highlight"); h.set(qn("w:val"), "yellow"); rpr.append(h)
    if align is not None:
        p.alignment = align


def remove_par(p):
    p._p.getparent().remove(p._p)


# annotation paragraphs like "<Times New Roman, Font 12, Italic>" (keep paragraphs that carry a section break)
for p in list(doc.paragraphs):
    t = p.text.strip()
    if (re.match(r"^<\s*Times New Roman", t) and "November" not in t) or t.startswith("(IF TEAM PROJECT"):
        if has_sect(p._p):
            set_par_text(p, [("", False)])
        else:
            remove_par(p)

# cover and title pages (two copies)
for k in range(2):
    p = [x for x in doc.paragraphs if x.text.strip() == "TITLE OF THE PROJECT REPORT"][0]
    set_par_text(p, [(TITLE.upper(), False)], size=20, bold=True)
    p = [x for x in doc.paragraphs if x.text.startswith("Bachelor of Technology in")][k]
    set_par_text(p, [("Bachelor of Technology in ", False), ("[PROGRAM NAME WITH SPECIALIZATION]", True)], size=22, bold=True)
    p = [x for x in doc.paragraphs if x.text.strip().startswith("NAME OF THE CANDIDATE")][0]
    set_par_text(p, [("NAREN MEDARAMETLA (", False), ("[REG. NO.]", True), (")", False)], size=16, bold=True)
    p.paragraph_format.space_after = Pt(2)
    new = copy.deepcopy(p._p)
    p._p.addnext(new)
    from docx.text.paragraph import Paragraph
    set_par_text(Paragraph(new, p._parent), [("RAMINENI DURGA SAI SUMANTH (", False), ("[REG. NO.]", True), (")", False)], size=16, bold=True)
for p in doc.paragraphs:
    if "November" in p.text:
        set_par_text(p, [("[MONTH, YEAR]", True)], size=12, bold=False)

# declaration
p = find_par("I hereby declare that the thesis entitled")
set_par_text(p, [("I hereby declare that the thesis entitled “", False), (TITLE, False), ("” submitted by ", False),
                 ("[YOUR NAME] ([REGISTER NUMBER])", True),
                 (", for the award of the degree of Bachelor of Technology in ", False), ("[PROGRAM NAME]", True),
                 (", Vellore Institute of Technology, Chennai is a record of bonafide work carried out by me under the supervision of Dr. Kowsigan M.", False)])
p = find_par("Date:") if find_par("Signature of the Candidate") is None else find_par("Signature of the Candidate")
for r in p.runs:
    if "Date:" in r.text:
        r.text = r.text.replace("Date:", "Date: ")
date_run = p.add_run("")
# put the date placeholder right after "Date:"
txt = p.text
set_par_text(p, [("Date: ", False), ("[DATE]", True), ("\t" + txt.split("Date:")[1].strip().lstrip("\t"), False)])

# certificate
p = find_par("This is to certify that the report entitled")
set_par_text(p, [("This is to certify that the report entitled “", False), (TITLE, False), ("” is prepared and submitted by ", False),
                 ("[YOUR NAME] ([REGISTER NUMBER])", True),
                 (" to Vellore Institute of Technology, Chennai, in partial fulfillment of the requirement for the award of the degree of Bachelor of Technology in ", False),
                 ("[PROGRAM NAME]", True),
                 (" is a bonafide record carried out under my guidance. The project fulfills the requirements as per the regulations of this University and in my opinion meets the necessary standards for submission. The contents of this report have not been submitted and will not be submitted either in part or in full, for the award of any other degree or diploma and the same is certified.", False)])
p = find_par("Name: Dr./Prof.")
set_par_text(p, [("Name: Dr. Kowsigan M", False)])
p = find_par("(MENTION PROGRAM NAME")
set_par_text(p, [("[PROGRAM NAME WITH SPECIALIZATION]", True)])
p = find_par("Name:(HoD NAME)")
set_par_text(p, [("           Name: ", False), ("[HoD NAME]", True)])
p = find_par("(Seal of SCHOOL")
set_par_text(p, [("                                         (Seal of the School of Computer Science and Engineering)", False)])

# abstract
ABSTRACT = [
    "Deepfake detectors are usually trained once and then deployed, but new face manipulation tools appear every few months, so a deployed detector has to be updated with the few examples of a new forgery that its operator can collect. This project studies what such updates cost and how to reduce the cost without keeping any data from earlier training. We start from GenD, a recent detector that adapts only the layer-normalization parameters of a frozen CLIP ViT-L/14 image encoder and generalizes well to forgeries it was never trained on. Updating it sequentially on few-shot samples of DFDC-Preview, DeepFakeDetection and Celeb-DF improves performance on those datasets and causes almost no conventional forgetting, but it lowers detection of manipulations outside the stream and shifts the decision threshold of the original task.",
    "We adapt genuine-sample alignment, a regularizer proposed for continual fake-audio detection, which keeps the previous model's features of genuine faces while the features of manipulated faces are free to change, and combine it with logit distillation. On two sealed datasets that no model saw during training or model selection, WildDeepfake and DeepFakeFace, the combined update keeps the unseen-forgery AUROC within 0.19 points of the never-updated detector on average over five training seeds, 0.99 points above learning without forgetting and 2.58 points above plain fine-tuning, and it outperforms a re-implementation of the GPL method in every setting tested. The advantage holds across three task orders, a second foundation encoder and image degradations, persists when competitors are tuned to learn as much, and grows over nine updates. Control experiments show that the benefit depends on the genuine class, and a ranking-margin analysis explains why. Protection is incomplete: text-to-image faces still lose 4.7 points, new datasets are learned less completely, and decision thresholds do not transfer to unseen data.",
]
ap = find_par("Electronic  Business")
set_par_text(ap, [(ABSTRACT[0], False)])
ap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
ap.paragraph_format.line_spacing = 1.5
ap.paragraph_format.first_line_indent = Inches(0.5)
ap.paragraph_format.space_after = Pt(6)
new = copy.deepcopy(ap._p); ap._p.addnext(new)
from docx.text.paragraph import Paragraph  # noqa: E402
p2 = Paragraph(new, ap._parent)
set_par_text(p2, [(ABSTRACT[1], False)])
kw = copy.deepcopy(new); new.addnext(kw)
p3 = Paragraph(kw, ap._parent)
set_par_text(p3, [("Keywords: ", False), ("deepfake detection, continual learning, knowledge distillation, foundation models, generalization, face forgery", False)])
p3.runs[0].font.bold = True
p3.paragraph_format.first_line_indent = Inches(0)
p3.paragraph_format.space_before = Pt(6)
# remove the empty paragraphs after the abstract (up to the section-break paragraph) so the page does not overflow
el = kw.getnext()
while el is not None and el.tag == qn("w:p") and not has_sect(el) and not ptext(el).strip() and el.find(".//" + qn("w:drawing")) is None:
    nxt = el.getnext(); body.remove(el); el = nxt
if el is not None and has_sect(el):
    for t in el.iter(qn("w:t")):
        t.text = ""

# acknowledgement
p = find_par("It is my pleasure to express")
set_par_text(p, [("It is my pleasure to express with deep sense of gratitude to Dr. Kowsigan M,", False)])
p = find_par("<Designation>")
set_par_text(p, [("Associate Professor, School of Computer Science and Engineering, Vellore Institute of Technology, Chennai, for his constant guidance, continual encouragement and understanding; more than all, he taught me patience in my endeavor. My association with him is not confined to academics only, but it is a great opportunity on my part to work with an intellectual and expert in the field of machine learning.", False)])
p = find_par("In jubilant state")
set_par_text(p, [("I express my sincere thanks to ", False), ("[HoD NAME]", True), (", Head of the Department, B.Tech. ", False),
                 ("[PROGRAM NAME]", True), (", and the Project Coordinators for their valuable support and encouragement to take up and complete the thesis.", False)])
p = find_par("Name of the student")
txt = p.text
set_par_text(p, [("Date: ", False), ("[DATE]", True), ("\t\t\t\t\t\t", False), ("[YOUR NAME]", True)])
for q in doc.paragraphs:
    for r in q.runs:
        if r.text.strip() and r._r.find(qn("w:rPr")) is not None and r._r.find(qn("w:rPr")).find(qn("w:highlight")) is not None \
                and not r.text.startswith("["):
            r._r.find(qn("w:rPr")).remove(r._r.find(qn("w:rPr")).find(qn("w:highlight")))
        if r.text.strip() and r.font.color is not None and r.font.color.rgb is not None and str(r.font.color.rgb) not in ("000000",):
            r.font.color.rgb = RGBColor(0, 0, 0)
# ---- layout fixes for the A4 page: signature lines with a right tab, fewer empty lines on the certificate
def right_tab_line(p, left_segments, right_segments):
    ppr = p._p.get_or_add_pPr()
    for t in ppr.findall(qn("w:tabs")):
        ppr.remove(t)
    tabs_el = OxmlElement("w:tabs"); t = OxmlElement("w:tab"); t.set(qn("w:val"), "right"); t.set(qn("w:pos"), str(int(TEXT_W * 1440)))
    tabs_el.append(t); ppr.append(tabs_el)
    p.paragraph_format.first_line_indent = Inches(0); p.paragraph_format.left_indent = Inches(0)
    set_par_text(p, left_segments + [("\t", False)] + right_segments)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT


right_tab_line(find_par("Signature of the Candidate"), [("Date: ", False), ("[DATE]", True)], [("Signature of the Candidate", False)])
right_tab_line(find_par("[YOUR NAME]", 2) or [q for q in doc.paragraphs if q.text.startswith("Date: [DATE]")][-1],
               [("Date: ", False), ("[DATE]", True)], [("[YOUR NAME]", True)])

kids = list(body.iterchildren())
ci = next(i for i, e in enumerate(kids) if ptext(e).strip() == "CERTIFICATE")
ce = next(i for i in range(ci, len(kids)) if kids[i].tag == qn("w:p") and has_sect(kids[i]))
prev_empty = False
for e in kids[ci + 1: ce]:
    if e.tag != qn("w:p"):
        continue
    empty = not ptext(e).strip() and e.find(".//" + qn("w:drawing")) is None
    if empty and prev_empty:
        body.remove(e)
    prev_empty = empty
for key in ("Approved by the Head of Department", "[PROGRAM NAME WITH SPECIALIZATION]", "Name: [HoD NAME]", "(Seal of the School"):
    hits = [x for x in doc.paragraphs if key in x.text]
    if not hits:
        print("MISSING KEY", key, [x.text for x in doc.paragraphs if "Seal" in x.text or "HoD" in x.text]); continue
    q = hits[-1]
    txt = q.text.strip()
    segs = [(txt, False)] if "[" not in txt else [(a, a.startswith("[")) for a in re.split(r"(\[[^\]]+\])", txt) if a]
    set_par_text(q, segs)
    q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    q.paragraph_format.left_indent = Inches(0); q.paragraph_format.first_line_indent = Inches(0)
for q in doc.paragraphs:
    if q.text.strip() == "Date:" and q.text.startswith("    "):
        set_par_text(q, [("Date:", False)]); q.alignment = WD_ALIGN_PARAGRAPH.CENTER

# acknowledgement: join the guide sentence that the template splits over two paragraphs, tighten spacing
a1 = find_par("It is my pleasure to express"); a2 = find_par("Associate Professor, School of Computer Science")
set_par_text(a1, [(a1.text.strip() + " " + a2.text.strip(), False)])
remove_par(a2)
kids = list(body.iterchildren())
ai = next(i for i, e in enumerate(kids) if ptext(e).strip() == "ACKNOWLEDGEMENT")
ae = next(i for i in range(ai, len(kids)) if kids[i].tag == qn("w:p") and has_sect(kids[i]))
for e in kids[ai + 1: ae]:
    if e.tag == qn("w:p") and not ptext(e).strip():
        body.remove(e)
for q in doc.paragraphs:
    if q.text.startswith(("It is my pleasure", "It is with gratitude", "Special mention", "I express my sincere", "My sincere thanks")):
        q.paragraph_format.space_before = Pt(6); q.paragraph_format.line_spacing = 1.15
        q.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        set_par_text(q, [(re.sub(r"\s{2,}", " ", q.text.strip()), False)] if "[" not in q.text else
                     [(a, a.startswith("[")) for a in re.split(r"(\[[^\]]+\])", re.sub(r"\s{2,}", " ", q.text.strip())) if a])

# manual page numbers typed in the template ("i", "ii") are replaced by footer fields
for p in doc.paragraphs:
    if p.text.strip() in ("i", "ii") and has_sect(p._p):
        for r in p.runs:
            r.text = ""

# ======================================================================================= generated content
def add_heading1(lines_sizes, label_for_toc=None):
    p = doc.add_paragraph(style="Heading 1")
    for i, (txt, size) in enumerate(lines_sizes):
        if i:
            p.add_run().add_break()
        add_run(p, txt, size, bold=True)
    return p


def add_heading(level, num, title):
    p = doc.add_paragraph(style=f"Heading {level}")
    add_run(p, f"{num}  ", 14 if level == 2 else 12)
    render_inline(p, title, 14 if level == 2 else 12)
    return p


def caption_par(style, label, num, text, size=11):
    p = doc.add_paragraph(style=style)
    add_run(p, f"{label} {num}: ", size, bold=True)
    render_inline(p, text, size)
    return p


def add_figure(d):
    p = doc.add_paragraph(style="Figure Para")
    p.add_run().add_picture(os.path.join(HERE, d["path"]), width=Inches(min(d["width"], TEXT_W)))
    caption_par("Fig Caption", "Figure", d["num"], d["caption"])
    figs.append((d["num"], resolve_refs(d["caption"])))


def add_table(d, keep_last=False):
    caption_par("Tab Caption", "Table", d["num"], d["caption"])
    tabs.append((d["num"], resolve_refs(d["caption"])))
    rows, widths, size = d["rows"], d["widths"], d["size"]
    nc = len(widths)
    for r in rows:
        if len(r) != nc:
            raise ValueError(f"table {d['label']}: row has {len(r)} cells, expected {nc}: {r}")
    if sum(widths) > TEXT_W:
        widths = [w * TEXT_W / sum(widths) for w in widths]
    t = doc.add_table(rows=len(rows), cols=nc)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    no_table_borders(t)
    set_widths(t, widths)
    cell_margins(t)
    align = d["align"] or ("l" + "c" * (nc - 1))
    for i, r in enumerate(rows):
        tr = t.rows[i]._tr
        trPr = tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
        if i == 0:
            trPr.append(OxmlElement("w:tblHeader"))
        for j, txt in enumerate(r):
            c = t.cell(i, j)
            p = c.paragraphs[0]
            p.style = doc.styles["Table Text"]
            p.alignment = {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER, "r": WD_ALIGN_PARAGRAPH.RIGHT}[align[j]]
            p.paragraph_format.space_before = Pt(1.5); p.paragraph_format.space_after = Pt(1.5)
            if i < len(rows) - 1 or keep_last:
                p.paragraph_format.keep_with_next = True
            small = d.get("intervals") and i > 0 and not r[0]
            if small:
                p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(2)
            elif d.get("intervals") and i > 0:
                p.paragraph_format.space_after = Pt(0)
            render_inline(p, txt, size - 1.5 if small else size, bold=(i == 0))
            if i == 0:
                set_cell_borders(c, top=("single", 12), bottom=("single", 6))
            if i == len(rows) - 1:
                set_cell_borders(c, bottom=("single", 12))
    return t


def add_note(text):
    p = doc.add_paragraph(style="Table Note")
    render_inline(p, text, 10)


def add_equation(d):
    t = doc.add_table(rows=1, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    no_table_borders(t)
    set_widths(t, [0.7, TEXT_W - 1.4, 0.7])
    cell_margins(t, 0, 0)
    mid, right = t.cell(0, 1), t.cell(0, 2)
    for c in (t.cell(0, 0), mid, right):
        tcPr = c._tc.get_or_add_tcPr()
        va = OxmlElement("w:vAlign"); va.set(qn("w:val"), "center"); tcPr.append(va)
    p = mid.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(4)
    omp = etree.SubElement(p._p, "{%s}oMathPara" % MATH_NS)
    omp.append(math_element("DISPLAY::" + d["tex"], 12))
    q = right.paragraphs[0]
    q.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_run(q, f"({d['num']})", 12)


def add_algorithm(d):
    t = doc.add_table(rows=2, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    no_table_borders(t)
    set_widths(t, [TEXT_W])
    cell_margins(t, 80, 40)
    c0, c1 = t.cell(0, 0), t.cell(1, 0)
    set_cell_borders(c0, top=("single", 12), bottom=("single", 6))
    set_cell_borders(c1, bottom=("single", 12))
    for tr in t.rows:
        tr._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    p = c0.paragraphs[0]
    p.paragraph_format.keep_with_next = True
    add_run(p, f"Algorithm {d['num']}: ", 11, bold=True)
    render_inline(p, d["caption"], 11)
    cell = c1
    first = True

    def line_par():
        nonlocal first
        if first:
            first = False
            return cell.paragraphs[0]
        return cell.add_paragraph()

    if d["require"]:
        p = line_par()
        p.paragraph_format.left_indent = Inches(0.35); p.paragraph_format.first_line_indent = Inches(-0.35)
        p.paragraph_format.space_after = Pt(2)
        add_run(p, "Require: ", 11, bold=True)
        render_inline(p, d["require"], 11)
    for n, (lvl, txt) in enumerate(d["steps"], 1):
        p = line_par()
        ind = 0.35 + 0.3 * lvl
        p.paragraph_format.left_indent = Inches(ind)
        p.paragraph_format.first_line_indent = Inches(-ind)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(ind))
        add_run(p, f"{n}:", 10)
        add_run(p, "\t", 11)
        main, _, comment = txt.partition("//")
        render_inline(p, main.strip(), 11)
        if comment:
            add_run(p, "   ▷ ", 10, color="555555")
            render_inline(p, comment.strip(), 10, italic=True)
    for pp in cell.paragraphs:
        pp.paragraph_format.line_spacing = 1.1
    sp = doc.add_paragraph(style="Report Plain"); sp.paragraph_format.space_after = Pt(0)


def add_code(code):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    no_table_borders(t)
    set_widths(t, [TEXT_W])
    cell_margins(t, 100, 60)
    c = t.cell(0, 0)
    tcPr = c._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "F2F2F2")
    tcPr.append(shd)
    set_cell_borders(c, top=("single", 4), bottom=("single", 4), left=("single", 4), right=("single", 4))
    lines = code.split("\n")
    for i, ln in enumerate(lines):
        p = c.paragraphs[0] if i == 0 else c.add_paragraph()
        p.style = doc.styles["Code Text"]
        r = add_run(p, ln if ln else " ", 8.5, font="Courier New")
        for tnode in r._r.findall(qn("w:t")):
            tnode.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    sp = doc.add_paragraph(style="Report Plain"); sp.paragraph_format.space_after = Pt(0)


def emit(blocks, chap_label):
    for bi, (kind, d) in enumerate(blocks):
        nxt = blocks[bi + 1][0] if bi + 1 < len(blocks) else None
        if kind == "chapter":
            p = add_heading1([(chap_label, 12), (d["title"], 16)])
            heads.append((1, chap_label, d["title"]))
        elif kind == "appendix":
            p = add_heading1([(f"APPENDIX {d['num']}", 12), (d["title"], 16)])
            heads.append((1, f"APPENDIX {d['num']}", d["title"]))
        elif kind == "section":
            add_heading(2, d["num"], d["title"])
            heads.append((2, d["num"], d["title"]))
        elif kind == "subsection":
            add_heading(3, d["num"], d["title"])
            heads.append((3, d["num"], d["title"]))
        elif kind == "p":
            body_par(d)
        elif kind == "note":
            add_note(d)
        elif kind == "fig":
            add_figure(d)
        elif kind == "table":
            add_table(d, keep_last=(nxt == "note"))
            if nxt != "note":
                sp = doc.add_paragraph(style="Table Note"); sp.paragraph_format.space_after = Pt(4)
        elif kind == "eq":
            add_equation(d)
        elif kind == "alg":
            add_algorithm(d)
        elif kind == "code":
            add_code(d)


# --- front lists: contents, figures, tables, symbols and abbreviations (filled after the body is generated)
def field_paragraphs(instr, entries, style_for):
    """Insert a TOC-type field whose cached result lists the entries [(level, text, page)]."""
    ps = []
    for k, (lvl, text, page) in enumerate(entries):
        p = doc.add_paragraph(style=style_for(lvl))
        if k == 0:
            r = OxmlElement("w:r"); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "begin"); r.append(fc); p._p.append(r)
            r = OxmlElement("w:r"); it = OxmlElement("w:instrText"); it.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            it.text = f" {instr} "; r.append(it); p._p.append(r)
            r = OxmlElement("w:r"); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "separate"); r.append(fc); p._p.append(r)
        render_inline(p, text, 12, bold=(lvl == 1 and style_for(lvl) == "toc 1"))
        add_run(p, "\t" + str(page), 12, bold=(lvl == 1 and style_for(lvl) == "toc 1"))
        ps.append(p)
    r = OxmlElement("w:r"); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "end"); r.append(fc); ps[-1]._p.append(r)


def front_heading(text):
    p = doc.add_paragraph(style="Heading 1")
    add_run(p, text, 16, bold=True)
    return p


# section break paragraph after the acknowledgement is already in the template; the front lists start a new section
FRONT_START = len(list(body.iterchildren()))

front_heading("CONTENTS")
toc_placeholder = doc.add_paragraph()  # replaced later
front_heading("LIST OF TABLES")
lot_placeholder = doc.add_paragraph()
front_heading("LIST OF FIGURES")
lof_placeholder = doc.add_paragraph()
front_heading("LIST OF SYMBOLS AND ABBREVIATIONS")
ABBR = [
    ("AA", "Average AUROC over all trained tasks after the final update"),
    ("ACC", "Accuracy at a fixed decision threshold"),
    ("AF", "Average forgetting"),
    ("AUROC", "Area under the receiver operating characteristic curve"),
    ("CDF", "Celeb-DF dataset"),
    ("CE", "Cross-entropy"),
    ("CLIP", "Contrastive Language-Image Pre-training"),
    ("CUDA", "Compute Unified Device Architecture"),
    ("DF, F2F, FS, NT", "Deepfakes, Face2Face, FaceSwap and NeuralTextures (FaceForensics++ methods)"),
    ("DFD", "DeepFakeDetection dataset"),
    ("DFDC-P", "Deepfake Detection Challenge preview dataset"),
    ("DFF", "DeepFakeFace dataset"),
    ("DFWF", "Detecting Fake Without Forgetting"),
    ("ECE", "Expected calibration error"),
    ("EWC", "Elastic weight consolidation"),
    ("FF++", "FaceForensics++ dataset"),
    ("FPR, TPR", "False-positive rate, true-positive rate"),
    ("FSh", "FaceShifter manipulation"),
    ("FT", "Fine-tuning"),
    ("GA", "Genuine-sample alignment"),
    ("GA+LwF", "Genuine-sample alignment combined with learning without forgetting"),
    ("GAN", "Generative adversarial network"),
    ("GenD", "Generalizable deepfake detector of Yermakov et al. (2026)"),
    ("GPL", "Generalization-preserved learning"),
    ("GPU", "Graphics processing unit"),
    ("KD", "Knowledge distillation"),
    ("KL", "Kullback-Leibler divergence"),
    ("LN", "Layer normalization"),
    ("LwF", "Learning without forgetting"),
    ("MLP", "Multilayer perceptron"),
    ("PE", "Perception Encoder"),
    ("ROC", "Receiver operating characteristic"),
    ("SDC", "Semantic drift compensation"),
    ("Seq. FT", "Sequential fine-tuning"),
    ("T2I", "Text-to-image"),
    ("t-SNE", "t-distributed stochastic neighbor embedding"),
    ("ViT", "Vision transformer"),
    ("WDF", "WildDeepfake dataset"),
]


def two_col(rows, w=(1.5, 4.2), math_first=False):
    t = doc.add_table(rows=len(rows), cols=2)
    no_table_borders(t)
    set_widths(t, list(w))
    cell_margins(t, 40, 10)
    for i, (a, b) in enumerate(rows):
        pa, pb = t.cell(i, 0).paragraphs[0], t.cell(i, 1).paragraphs[0]
        for pp in (pa, pb):
            pp.paragraph_format.space_after = Pt(3); pp.paragraph_format.line_spacing = 1.15
        if math_first:
            pa._p.append(math_element(a, 12))
        else:
            add_run(pa, a, 12)
        render_inline(pb, b, 12)


p = doc.add_paragraph(style="Report Plain"); add_run(p, "Abbreviations", 12, bold=True)
two_col(ABBR)
p = doc.add_paragraph(style="Report Plain"); p.paragraph_format.space_before = Pt(12); add_run(p, "Symbols", 12, bold=True)
two_col(SYMBOLS, math_first=True)

# section break: front matter (roman) ends here
sb = doc.add_paragraph()
sect_front = OxmlElement("w:sectPr")
sb._p.get_or_add_pPr().append(sect_front)

# --- chapters
for n, blocks in enumerate(chapters, 1):
    emit(blocks, f"CHAPTER {n}")

# --- references
front_heading("REFERENCES")
heads.append((1, "", "REFERENCES"))
keys = sorted(USED_CITES | set(), key=sort_key)
for i, k in enumerate(keys, 1):
    p = doc.add_paragraph(style="Reference Entry")
    render_inline(p, f"[{i}] " + apa_entry(k), 12)
unused = set(REFS) - USED_CITES
if unused:
    print("WARNING: references not cited:", sorted(unused))

# --- appendices
emit(appendix, None)

# ======================================================================================= front lists content
def page_of(key, default="0"):
    return str(PAGES.get(key, default))


toc_entries = [(1, h, page_of("H|" + h)) for h in ("CONTENTS", "LIST OF TABLES", "LIST OF FIGURES", "LIST OF SYMBOLS AND ABBREVIATIONS")]
for lvl, num, title in heads:
    if lvl == 1:
        text = f"{num}  {title}".strip()
    else:
        text = f"{num}  {title.upper()}"
    toc_entries.append((lvl, text, page_of("H|" + (num + " " + title).strip())))
lof_entries = [(9, f"Figure {n}: {c}", page_of("F|" + n)) for n, c in figs]
lot_entries = [(9, f"Table {n}: {c}", page_of("T|" + n)) for n, c in tabs]

# build the field paragraphs at the end of the document, then move them into place
def build_and_move(placeholder, instr, entries, style_for):
    before = len(list(body.iterchildren()))
    field_paragraphs(instr, entries, style_for)
    kids = list(body.iterchildren())
    new = kids[before - 1: -1] if kids[-1].tag == qn("w:sectPr") else kids[before:]
    # the final sectPr is the last child; new paragraphs were inserted before it
    new = [el for el in kids if el.tag == qn("w:p")][-len(entries):]
    anchor = placeholder._p
    for el in new:
        anchor.addnext(el); anchor = el
    placeholder._p.getparent().remove(placeholder._p)


build_and_move(toc_placeholder, 'TOC \\o "1-3" \\h \\z \\u', toc_entries, lambda l: {1: "toc 1", 2: "toc 2", 3: "toc 3"}[l])
build_and_move(lof_placeholder, 'TOC \\h \\z \\t "Fig Caption,9"', lof_entries, lambda l: "toc 9")
build_and_move(lot_placeholder, 'TOC \\h \\z \\t "Tab Caption,9"', lot_entries, lambda l: "toc 9")

# ======================================================================================= sections, page size, numbering
def footer_with_page(section, numbered):
    section.footer.is_linked_to_previous = False
    section.header.is_linked_to_previous = False
    for p in section.header.paragraphs:
        for r in p.runs:
            r.text = ""
    ft = section.footer
    for p in ft.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    p = ft.paragraphs[0]
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    for child in list(p._p):
        if child.tag != qn("w:pPr"):
            p._p.remove(child)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if numbered:
        for kind, text in (("begin", None), ("instr", " PAGE "), ("separate", None), ("text", "1"), ("end", None)):
            r = OxmlElement("w:r")
            rpr = OxmlElement("w:rPr"); rf = OxmlElement("w:rFonts")
            for a in ("ascii", "hAnsi"):
                rf.set(qn("w:" + a), "Times New Roman")
            rpr.append(rf); sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "24"); rpr.append(sz); r.append(rpr)
            if kind in ("begin", "separate", "end"):
                fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); r.append(fc)
            elif kind == "instr":
                it = OxmlElement("w:instrText"); it.set("{http://www.w3.org/XML/1998/namespace}space", "preserve"); it.text = text; r.append(it)
            else:
                t = OxmlElement("w:t"); t.text = text; r.append(t)
            p._p.append(r)


def set_page(sectPr, fmt=None, start=None):
    for e in sectPr.findall(qn("w:titlePg")):
        sectPr.remove(e)
    for tag in ("w:headerReference", "w:footerReference"):
        for e in sectPr.findall(qn(tag)):
            if e.get(qn("w:type")) in ("first", "even"):
                sectPr.remove(e)
    for tag in ("w:pgSz", "w:pgMar", "w:pgNumType", "w:cols"):
        for e in sectPr.findall(qn(tag)):
            sectPr.remove(e)
    sz = OxmlElement("w:pgSz"); sz.set(qn("w:w"), "11906"); sz.set(qn("w:h"), "16838")
    mar = OxmlElement("w:pgMar")
    for k, v in (("top", 1440), ("right", 1440), ("bottom", 1440), ("left", 2160), ("header", 708), ("footer", 708), ("gutter", 0)):
        mar.set(qn("w:" + k), str(v))
    # schema order: footnotePr/endnotePr/type before pgSz; insert after any of those
    pos = 0
    for i, ch in enumerate(sectPr):
        if ch.tag in (qn("w:headerReference"), qn("w:footerReference"), qn("w:footnotePr"), qn("w:endnotePr"), qn("w:type")):
            pos = i + 1
    sectPr.insert(pos, sz); sectPr.insert(pos + 1, mar)
    if fmt:
        pn = OxmlElement("w:pgNumType"); pn.set(qn("w:fmt"), fmt)
        if start is not None:
            pn.set(qn("w:start"), str(start))
        # pgNumType goes after pgMar/paperSrc/pgBorders/lnNumType
        idx = list(sectPr).index(mar) + 1
        while idx < len(sectPr) and sectPr[idx].tag in (qn("w:paperSrc"), qn("w:pgBorders"), qn("w:lnNumType")):
            idx += 1
        sectPr.insert(idx, pn)
    cols = OxmlElement("w:cols"); cols.set(qn("w:space"), "720")
    idx = list(sectPr).index(sectPr.find(qn("w:pgNumType")) if sectPr.find(qn("w:pgNumType")) is not None else mar) + 1
    sectPr.insert(idx, cols)


sections = doc.sections
print("sections:", len(sections))
# expected: cover, title, declaration+certificate, abstract, acknowledgement, front lists, main
assert len(sections) == 7, len(sections)
plan = [(False, None, None), (False, None, None), (False, None, None), (True, "lowerRoman", 1), (True, "lowerRoman", None),
        (True, "lowerRoman", None), (True, "decimal", 1)]
for s, (numbered, fmt, start) in zip(sections, plan):
    footer_with_page(s, numbered)
    set_page(s._sectPr, fmt, start)

# update fields when the document is opened in Word (updateFields must precede these settings elements)
settings = doc.settings.element
uf = settings.find(qn("w:updateFields"))
if uf is None:
    uf = OxmlElement("w:updateFields")
    later = ("hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr", "attachedSchema",
             "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats", "doNotAutoCompressPictures", "forceUpgrade",
             "captions", "readModeInkLockDown", "smartTagType", "schemaLibrary", "shapeDefaults", "doNotEmbedSmartTags",
             "decimalSymbol", "listSeparator")
    pos = next((i for i, ch in enumerate(settings) if etree.QName(ch).localname in later), len(settings))
    settings.insert(pos, uf)
uf.set(qn("w:val"), "true")

# ---- put the children of property elements into schema order (Word rejects out-of-order elements)
ORDER = {
    "pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange",
    "rPr": "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs em lang eastAsianLayout specVanish oMath",
    "tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing tblInd tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription",
    "trPr": "cnfStyle divId gridBefore gridAfter wBefore wAfter cantSplit trHeight tblHeader tblCellSpacing jc hidden",
    "tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark",
    "sectPr": "headerReference footerReference footnotePr endnotePr type pgSz pgMar paperSrc pgBorders lnNumType pgNumType cols formProt vAlign noEndnote titlePg textDirection bidi rtlGutter docGrid printerSettings sectPrChange",
    "tcBorders": "top start left bottom end right insideH insideV tl2br tr2bl",
    "tblBorders": "top start left bottom end right insideH insideV",
    "tblCellMar": "top start left bottom end right",
}
ORDER = {k: {t: i for i, t in enumerate(v.split())} for k, v in ORDER.items()}


def reorder(root):
    n = 0
    for el in root.iter():
        if not isinstance(el.tag, str) or not el.tag.startswith("{%s}" % W_NS):
            continue
        name = etree.QName(el).localname
        if name not in ORDER or (name == "rPr" and el.getparent() is not None and etree.QName(el.getparent()).namespace == MATH_NS):
            continue
        rank = ORDER[name]
        kids = list(el)
        keyed = sorted(range(len(kids)), key=lambda i: (rank.get(etree.QName(kids[i]).localname, 1000) if etree.QName(kids[i]).namespace == W_NS else 1000, i))
        if keyed != list(range(len(kids))):
            for k in kids:
                el.remove(k)
            for i in keyed:
                el.append(kids[i])
            n += 1
    return n


fixed = reorder(doc.element)
for s_ in doc.sections:
    for part in (s_.footer, s_.header):
        fixed += reorder(part._element)
fixed += reorder(doc.styles.element)
print("reordered elements:", fixed)

# document properties
cp = doc.core_properties
cp.title = TITLE
cp.subject = "Project report"
cp.author = ""
cp.last_modified_by = ""
cp.comments = ""

doc.save(OUT)
json.dump({"heads": heads, "figs": figs, "tabs": tabs}, open(os.path.join(HERE, "_structure.json"), "w"), indent=0)
print("saved", OUT, "figures", len(figs), "tables", len(tabs), "references", len(keys))
